"""Word (.docx) files of Discovery 1 and Discovery 2 versions (BAP-7): customer documents.

Server-side files, so both Discovery documents can be downloaded and attached to the follow-up
email like the presentations. Rendered from the stored version only - nothing is generated here.

The files go to the client. They therefore show the content and nothing about how it was made:
no item origin tags, no version or document ids, no transcript file names or revisions, no
generator names, and no item that rests only on the owner's personal notes unless the owner
shared it (``services.meeting.discovery2.is_internal_only``). Where a point comes from is told
in the client's terms - "from our conversation", "our observations", "from our earlier
analysis". The full provenance stays in the stored version and in the application.

BAP-13: every label is in the document's own language (``language`` of the stored version,
English when it has none), and a German file is marked as German for spelling and hyphenation.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from services.language import stored_language
from services.meeting.discovery2 import is_internal_only

NAVY = RGBColor(0x0D, 0x12, 0x40)
SLATE = RGBColor(0x51, 0x5C, 0x70)
ACCENT = RGBColor(0xDD, 0x3D, 0x00)
FONT = "Calibri"
_LOGO = Path(__file__).resolve().parents[4] / "borek_assets" / "logo_dark.png"
AUTHOR = "Borek Solutions Group"

_MONTHS = {
    "en": ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"),
    "de": ("Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"),
}

_LABELS: dict[str, dict[str, Any]] = {
    "en": {
        "lang": None,
        "version": "Version {number}",
        "draft": "Draft · not approved",
        "page": "Page",
        "d2_eyebrow": "After our first conversation",
        "d2_subtitle": "Discovery after our first conversation",
        "summary": "Summary",
        "starting_point": "Our starting point",
        "added": "Our first conversation added {counts}.",
        "added_none": "Our first conversation confirmed the starting point; it added no new points.",
        "and": " and ",
        "count_names": {
            "context": ("point of context", "points of context"),
            "requirements": ("requirement", "requirements"),
            "challenges": ("challenge", "challenges"),
            "priorities": ("priority", "priorities"),
            "opportunities": ("opportunity", "opportunities"),
            "discussed_solutions": ("solution we discussed", "solutions we discussed"),
            "decisions": ("decision", "decisions"),
            "follow_ups": ("next step", "next steps"),
            "open_questions": ("open question", "open questions"),
        },
        "groups": {"client": "From our conversation", "ours": "Our observations", "earlier": "From our earlier analysis"},
        "nothing": "Nothing recorded.",
        "closing_title": "Please correct us",
        "closing": (
            "This document summarises what we heard in our first conversation and what our earlier analysis "
            "established. If we misunderstood a point, please tell us: it is the basis for our next conversation."
        ),
        "core_thesis": "Core thesis",
        "opportunities": "Opportunities by area",
        "workflows": "Target workflows",
        "human_decision": "Human decision: {gate}",
        "conclusion": "Conclusion",
    },
    "de": {
        "lang": "de-DE",
        "version": "Version {number}",
        "draft": "Entwurf · nicht freigegeben",
        "page": "Seite",
        "d2_eyebrow": "Nach unserem ersten Gespräch",
        "d2_subtitle": "Discovery nach unserem ersten Gespräch",
        "summary": "Zusammenfassung",
        "starting_point": "Unser Ausgangspunkt",
        "added": "Unser erstes Gespräch hat {counts} ergeben.",
        "added_none": "Unser erstes Gespräch hat den Ausgangspunkt bestätigt; neue Punkte kamen nicht hinzu.",
        "and": " und ",
        "count_names": {
            "context": ("Punkt zum Kontext", "Punkte zum Kontext"),
            "requirements": ("Anforderung", "Anforderungen"),
            "challenges": ("Herausforderung", "Herausforderungen"),
            "priorities": ("Priorität", "Prioritäten"),
            "opportunities": ("Opportunity", "Opportunities"),
            "discussed_solutions": ("besprochene Lösung", "besprochene Lösungen"),
            "decisions": ("Entscheidung", "Entscheidungen"),
            "follow_ups": ("nächsten Schritt", "nächste Schritte"),
            "open_questions": ("offene Frage", "offene Fragen"),
        },
        "groups": {"client": "Aus unserem Gespräch", "ours": "Unsere Beobachtungen", "earlier": "Aus unserer vorherigen Analyse"},
        "nothing": "Nichts erfasst.",
        "closing_title": "Bitte korrigieren Sie uns",
        "closing": (
            "Dieses Dokument fasst zusammen, was wir in unserem ersten Gespräch gehört haben und was unsere vorherige "
            "Analyse ergeben hat. Wenn wir einen Punkt falsch verstanden haben, sagen Sie es uns bitte: Er ist die "
            "Grundlage für unser nächstes Gespräch."
        ),
        "core_thesis": "Kernthese",
        "opportunities": "Opportunities nach Bereich",
        "workflows": "Ziel-Workflows",
        "human_decision": "Menschliche Entscheidung: {gate}",
        "conclusion": "Fazit",
    },
}

# Said in the meeting: the transcript supports it.
_CLIENT_SOURCES = frozenset({"transcript", "both"})
_GROUPS = ("client", "ours", "earlier")


def _long_date(value: Any, language: str) -> str:
    text = str(value or "")
    if len(text) >= 10 and text[4] == "-" and text[7] == "-" and text[5:7].isdigit() and text[8:10].isdigit() and 1 <= int(text[5:7]) <= 12:
        day, month = int(text[8:10]), _MONTHS[language][int(text[5:7]) - 1]
        return f"{day}. {month} {text[0:4]}" if language == "de" else f"{day} {month} {text[0:4]}"
    return ""


def _version_line(version: dict[str, Any], labels: dict[str, Any], language: str) -> str:
    """"Version 2 · 27 October 2026" for an approved version; a draft says that it is not approved."""
    line = labels["version"].format(number=version.get("version_number"))
    if str(version.get("status") or "") != "approved":
        return f"{line} · {labels['draft']}"
    date = _long_date(version.get("approved_at"), language)
    return f"{line} · {date}" if date else line


def _save(document: Any) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _font(style: Any, *, size: float, bold: bool = False, color: RGBColor | None = None) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(key), FONT)


def _page_field(paragraph: Any) -> None:
    """A PAGE field: Word fills in the page number."""
    run = paragraph.add_run()
    run.font.size = Pt(8.5)
    run.font.color.rgb = SLATE
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def _small(paragraph: Any, text: str, *, size: float = 9.5, color: RGBColor = SLATE, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def _base(*, eyebrow: str, title: str, subtitle: str, version_line: str, footer: str, labels: dict[str, Any]) -> Any:
    """A Borek document: logo header, footer with page numbers, a title block and Borek colours."""
    document = Document()
    styles = document.styles
    _font(styles["Normal"], size=11, color=RGBColor(0x1F, 0x23, 0x3A))
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    _font(styles["Title"], size=26, bold=True, color=NAVY)
    _font(styles["Heading 1"], size=16, bold=True, color=NAVY)
    styles["Heading 1"].paragraph_format.space_before = Pt(18)
    _font(styles["Heading 2"], size=12.5, bold=True, color=NAVY)
    _font(styles["Heading 3"], size=11, bold=True, color=SLATE)
    if labels["lang"]:
        # Spelling, grammar and hyphenation in the document's language.
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
            lang = OxmlElement("w:lang")
            lang.set(qn("w:val"), labels["lang"])
            styles[name].element.get_or_add_rPr().append(lang)

    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2.3)
    section.top_margin, section.bottom_margin = Cm(2.6), Cm(2.2)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if _LOGO.is_file():
        header.add_run().add_picture(str(_LOGO), height=Cm(0.75))
    else:
        _small(header, "BOREK", size=11, color=NAVY, bold=True)
    foot = section.footer.paragraphs[0]
    _small(foot, f"{footer} · {labels['page']} ", size=8.5)
    _page_field(foot)

    # File properties: no tool names, no comments.
    properties = document.core_properties
    properties.author = AUTHOR
    properties.last_modified_by = AUTHOR
    properties.title = title
    properties.subject = subtitle
    properties.comments = ""
    properties.keywords = ""
    properties.category = ""
    properties.language = labels["lang"] or "en-GB"

    _small(document.add_paragraph(), eyebrow.upper(), color=ACCENT, bold=True)
    document.add_paragraph(title, style="Title")
    if subtitle and subtitle != title:
        _small(document.add_paragraph(), subtitle, size=13)
    _small(document.add_paragraph(), version_line)
    return document


def _group(item: dict[str, Any]) -> str:
    if item.get("origin") == "discovery1":
        return "earlier"
    if item.get("origin") == "meeting" and item.get("source") in _CLIENT_SOURCES:
        return "client"
    return "ours"  # shared notes of the owner and points the owner wrote


def customer_sections(paper: dict[str, Any]) -> list[dict[str, Any]]:
    """The Discovery 2 sections as the client sees them: without the internal-only items."""
    return [
        {"key": section["key"], "title": section["title"], "items": [item for item in section.get("items") or [] if not is_internal_only(item)]}
        for section in paper.get("sections") or []
    ]


def _counts(sections: list[dict[str, Any]], labels: dict[str, Any]) -> str:
    parts = []
    for section in sections:
        fresh = [item for item in section["items"] if item.get("origin") != "discovery1"]
        if fresh and section["key"] in labels["count_names"]:
            singular, plural = labels["count_names"][section["key"]]
            parts.append(f"{len(fresh)} {singular if len(fresh) == 1 else plural}")
    if len(parts) > 1:
        return ", ".join(parts[:-1]) + labels["and"] + parts[-1]
    return parts[0] if parts else ""


def render_discovery2_docx(paper: dict[str, Any], version: dict[str, Any]) -> bytes:
    """Discovery 2 for the client: summary, the sections in the client's terms, and an invitation to correct us."""
    language = stored_language(paper.get("language"))
    labels = _LABELS[language]
    client = str(paper.get("client_name") or "").strip()
    document = _base(
        eyebrow=labels["d2_eyebrow"],
        title=client or labels["d2_subtitle"],
        subtitle=labels["d2_subtitle"],
        version_line=_version_line(version, labels, language),
        footer=f"{client} · {labels['d2_subtitle']}" if client else labels["d2_subtitle"],
        labels=labels,
    )
    sections = customer_sections(paper)
    document.add_heading(labels["summary"], level=1)
    previous = str((paper.get("summary") or {}).get("previous") or "").strip()
    if previous:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{labels['starting_point']}: ").bold = True
        paragraph.add_run(previous)
    counts = _counts(sections, labels)
    document.add_paragraph(labels["added"].format(counts=counts) if counts else labels["added_none"])
    for section in sections:
        document.add_heading(str(section["title"]), level=1)
        if not section["items"]:
            document.add_paragraph(labels["nothing"])
            continue
        groups = {key: [item for item in section["items"] if _group(item) == key] for key in _GROUPS}
        shown = [key for key in _GROUPS if groups[key]]
        for key in shown:
            # A section that holds only the earlier analysis needs no sub-heading.
            if shown != ["earlier"]:
                document.add_heading(labels["groups"][key], level=3)
            for item in groups[key]:
                document.add_paragraph(str(item["text"]), style="List Bullet")
    document.add_heading(labels["closing_title"], level=1)
    document.add_paragraph(labels["closing"])
    return _save(document)


def render_discovery1_docx(paper: dict[str, Any], version: dict[str, Any]) -> bytes:
    """The approved Discovery analysis (schema 2.0) as a Word document."""
    language = stored_language(paper.get("language"))
    labels = _LABELS[language]
    analysis = paper["analysis"]
    framing = analysis["framing"]["document"]
    title = str(framing.get("title") or "Discovery")
    document = _base(
        eyebrow=str(framing.get("eyebrow") or framing.get("document_type") or "Discovery"),
        title=title,
        subtitle=str(framing.get("subtitle") or ""),
        version_line=_version_line(version, labels, language),
        footer=str(framing.get("short_title") or title),
        labels=labels,
    )
    document.add_heading(labels["core_thesis"], level=1)
    document.add_paragraph(str(analysis["research"]["core_thesis"]["text"]))
    document.add_heading(labels["opportunities"], level=1)
    for area in analysis.get("areas") or []:
        document.add_heading(str(area["name"]), level=2)
        if area.get("lead"):
            document.add_paragraph(str(area["lead"]))
        for opportunity in area.get("opportunities") or []:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{opportunity['title']}: ").bold = True
            paragraph.add_run(f"{opportunity.get('solution', '')} {opportunity.get('result', '')}".strip())
            for question in opportunity.get("discovery_questions") or []:
                document.add_paragraph(str(question), style="List Bullet 2")
    workflows = analysis.get("target_workflows") or []
    if workflows:
        document.add_heading(labels["workflows"], level=1)
        for workflow in workflows:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"{workflow['name']}: ").bold = True
            paragraph.add_run(f"{workflow.get('autonomous_flow', '')} {labels['human_decision'].format(gate=workflow.get('human_decision_gate', ''))}")
    closing = analysis.get("closing") or {}
    if closing.get("paragraphs"):
        document.add_heading(str(closing.get("headline") or closing.get("name") or labels["conclusion"]), level=1)
        for text in closing["paragraphs"]:
            document.add_paragraph(str(text))
    return _save(document)
