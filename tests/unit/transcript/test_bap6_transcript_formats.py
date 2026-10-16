"""BAP-6 — any meeting transcript is accepted within the technical limits.

No structure is required: time stamps, speaker names and finding labels are optional. Files are
refused only when they are not a supported format, empty, corrupt or too large.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document

from services.transcript.ingestion import (
    MAX_TRANSCRIPT_BYTES,
    MAX_TRANSCRIPT_CHARACTERS,
    TranscriptIngestionError,
    ingest_transcript,
    transcript_mime_allowed,
)
from services.transcript.speaker_turns import UNKNOWN_SPEAKER, split_speaker_turns

UNSTRUCTURED = (
    "We spend hours every week reconciling invoices by hand.\n"
    "It would help if the ERP could match them automatically.\n"
    "Let's meet again next Tuesday with finance."
)


def turns(name: str, content: bytes) -> list[tuple[str, str]]:
    return [(turn.speaker, turn.text) for turn in split_speaker_turns(name, content)]


def docx_bytes(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table):
            for cell, value in zip(row.cells, values):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ structure is optional


def test_plain_unstructured_text_without_labels_names_or_times_is_kept_whole() -> None:
    parsed = turns("meeting.txt", UNSTRUCTURED.encode())
    assert parsed == [(UNKNOWN_SPEAKER, UNSTRUCTURED.replace("\n", " "))]


def test_a_short_transcript_is_accepted() -> None:
    assert turns("short.txt", b"Ok, let's start with the pilot.") == [(UNKNOWN_SPEAKER, "Ok, let's start with the pilot.")]


def test_a_long_multi_speaker_transcript_keeps_every_turn_and_name() -> None:
    speakers = ["Anna", "Ben", "Chloe Martin", "D. O'Neil"]
    lines = [f"{speakers[index % 4]}: Point number {index} about the rollout." for index in range(2000)]
    parsed = turns("long.txt", "\n".join(lines).encode())
    assert len(parsed) == 2000
    assert parsed[0] == ("Anna", "Point number 0 about the rollout.")
    assert parsed[1999] == ("D. O'Neil", "Point number 1999 about the rollout.")


def test_a_mix_of_named_and_unnamed_lines_keeps_all_content() -> None:
    parsed = turns("mixed.txt", b"Some opening words before anyone is named.\nAnna: Hello.\nand Anna keeps talking")
    assert parsed == [(UNKNOWN_SPEAKER, "Some opening words before anyone is named."), ("Anna", "Hello. and Anna keeps talking")]


def test_finding_keywords_are_not_required() -> None:
    text = ingest_transcript("meeting.txt", UNSTRUCTURED.encode()).normalized_text
    for keyword in ("Requirement", "Challenge", "Decision", "Opportunity"):
        assert keyword not in text


# ------------------------------------------------------------------ caption formats


def test_vtt_with_voice_tags_keeps_speakers_and_text() -> None:
    vtt = b"WEBVTT\n\n00:00:01.000 --> 00:00:04.000\n<v Anna>We need single sign-on.\n\n00:00:05.000 --> 00:00:07.000\n<v Ben>Agreed.\n"
    assert turns("call.vtt", vtt) == [("Anna", "We need single sign-on."), ("Ben", "Agreed.")]


def test_vtt_short_time_stamps_without_hours_are_cue_times_not_text() -> None:
    vtt = b"WEBVTT\n\n00:01.000 --> 00:04.000\nAnna: we need SSO\n\n00:05.500 --> 00:06.000\nok\n"
    assert turns("call.vtt", vtt) == [("Anna", "we need SSO"), (UNKNOWN_SPEAKER, "ok")]
    assert "-->" not in ingest_transcript("call.vtt", vtt).normalized_text


def test_srt_with_and_without_speakers() -> None:
    srt = b"1\n00:00:01,000 --> 00:00:02,000\nAnna: Hello there\n\n2\n0:00:03,000 --> 0:00:04,500\nnice to meet you\n"
    assert turns("call.srt", srt) == [("Anna", "Hello there"), (UNKNOWN_SPEAKER, "nice to meet you")]


def test_uppercase_extensions_are_the_same_format() -> None:
    assert turns("MEETING.TXT", b"Anna: Hi") == [("Anna", "Hi")]
    assert turns("CALL.SRT", b"1\n00:00:01,000 --> 00:00:02,000\nHi\n") == [(UNKNOWN_SPEAKER, "Hi")]
    assert turns("NOTES.DOCX", docx_bytes("Anna: Hi")) == [("Anna", "Hi")]


# ------------------------------------------------------------------ Word documents and encodings


def test_docx_paragraphs_and_tables_are_read_in_order() -> None:
    content = docx_bytes("Anna: Welcome everyone.", table=[["Ben", "We still export reports by hand."]])
    text = ingest_transcript("meeting.docx", content).normalized_text
    assert text == "Anna: Welcome everyone.\nBen We still export reports by hand."


def test_utf16_text_from_windows_notepad_is_decoded() -> None:
    parsed = turns("meeting.txt", "Anna: Grüße aus München".encode("utf-16"))
    assert parsed == [("Anna", "Grüße aus München")]


def test_windows_1252_text_is_decoded() -> None:
    assert turns("meeting.txt", "Anna: Café meeting".encode("cp1252")) == [("Anna", "Café meeting")]


# ------------------------------------------------------------------ what is refused, and why


@pytest.mark.parametrize("name", ["meeting.pdf", "meeting.doc", "meeting.rtf", "meeting.mp3", "meeting"])
def test_unsupported_formats_are_refused(name: str) -> None:
    with pytest.raises(TranscriptIngestionError, match="Unsupported transcript format"):
        ingest_transcript(name, b"hello")


@pytest.mark.parametrize("name", ["empty.txt", "empty.vtt", "empty.srt", "empty.docx"])
def test_empty_files_are_refused(name: str) -> None:
    with pytest.raises(TranscriptIngestionError, match="empty"):
        ingest_transcript(name, b"")


@pytest.mark.parametrize("name,content", [("blank.txt", b"  \n\n \t\n"), ("cues.vtt", b"WEBVTT\n\n00:00:01.000 --> 00:00:02.000\n\n")])
def test_files_without_any_text_are_refused(name: str, content: bytes) -> None:
    with pytest.raises(TranscriptIngestionError, match="no transcript text"):
        ingest_transcript(name, content)


def test_binary_content_renamed_to_txt_is_refused() -> None:
    png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x10"
    with pytest.raises(TranscriptIngestionError, match="does not contain plain text|not valid text"):
        ingest_transcript("image.txt", png)


@pytest.mark.parametrize("content", [b"PK\x03\x04 broken archive", b"plain text, not a document"])
def test_corrupt_docx_is_refused(content: bytes) -> None:
    with pytest.raises(TranscriptIngestionError, match="could not be read"):
        ingest_transcript("meeting.docx", content)


def test_docx_that_unpacks_beyond_the_limit_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * (101 * 1024 * 1024))
    assert len(buffer.getvalue()) < MAX_TRANSCRIPT_BYTES
    with pytest.raises(TranscriptIngestionError, match="larger than 10 MB"):
        ingest_transcript("bomb.docx", buffer.getvalue())


def test_size_and_length_limits() -> None:
    with pytest.raises(TranscriptIngestionError, match="larger than 10 MB"):
        ingest_transcript("big.txt", b"a" * (MAX_TRANSCRIPT_BYTES + 1))
    with pytest.raises(TranscriptIngestionError, match="too long to analyse"):
        ingest_transcript("long.txt", b"word " * (MAX_TRANSCRIPT_CHARACTERS // 5 + 10))
    at_limit = ("x" * 99 + "\n") * (MAX_TRANSCRIPT_CHARACTERS // 100)
    assert len(ingest_transcript("ok.txt", at_limit.encode()).normalized_text) <= MAX_TRANSCRIPT_CHARACTERS


# ------------------------------------------------------------------ MIME types browsers send


@pytest.mark.parametrize(
    "extension,mime",
    [
        (".txt", "text/plain"), (".txt", "text/plain; charset=utf-8"), (".txt", "TEXT/PLAIN"), (".txt", None), (".txt", ""),
        (".vtt", "text/vtt"), (".vtt", "text/vtt; charset=UTF-8"), (".vtt", "text/x-vtt"),
        (".srt", "application/x-subrip"), (".srt", "text/srt"), (".srt", "application/x-srt"), (".srt", "text/plain"),
        (".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        (".docx", "application/zip"), (".docx", "application/x-zip-compressed"), (".docx", "application/msword"),
        (".docx", "application/octet-stream"), (".DOCX", "application/zip"),
    ],
)
def test_common_browser_mime_variants_are_accepted(extension: str, mime: str | None) -> None:
    assert transcript_mime_allowed(extension, mime)


@pytest.mark.parametrize(
    "extension,mime",
    [
        (".txt", "image/png"), (".txt", "application/pdf"), (".txt", "application/zip"),
        (".txt", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        (".vtt", "video/mp4"), (".docx", "text/plain"), (".docx", "application/pdf"), (".pdf", "application/pdf"),
    ],
)
def test_mime_types_of_other_file_kinds_are_refused(extension: str, mime: str) -> None:
    assert not transcript_mime_allowed(extension, mime)
