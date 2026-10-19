"""BAP-4: "Let's talk" is the last slide of every Master Presentation (V1 and V2).

The closing slide is the canonical slide the master's manifest marks with the role "closing" -
never a fixed position. These tests inspect the final order of the plan, the stored slide specs
and the exported PPTX (and the PDF/previews rendered from it), for any number of appendix slides.
"""

from __future__ import annotations

import copy
import io
import json
import zipfile
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from services.presentation.borek_deck.deck_plan import borek_slide_spec
from services.presentation.borek_deck.render import render_deck_bundle
from services.presentation.master_deck import assembly, registry, v2
from services.presentation.master_deck.plan import CANONICAL_LAYOUT, MASTER_V1, MASTER_V2, build_master_plan, generation_manifest, same_source
from tests.fixtures.master_render_double import install_master_render_double
from tests.unit.presentation.test_master_presentation import RICH, approved_row
from tests.unit.presentation.test_master_presentation_v2 import snapshot_for

CLOSING_TITLE = "Let’s talk"
CLOSING_TEXT = "Managing Partner & Co-CCO"  # contact on the canonical "Let's talk" slide


@pytest.fixture(autouse=True)
def _renderer(monkeypatch: pytest.MonkeyPatch) -> None:
    install_master_render_double(monkeypatch)


def specs_of(plan: dict[str, Any]) -> list[dict[str, Any]]:
    return [borek_slide_spec(planned) for planned in plan["slides"]]


def texts(slide: Any) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def exported(specs: list[dict[str, Any]], deck_kind: str, tmp_path: Path) -> tuple[Path, dict[str, Any], zipfile.ZipFile]:
    bundle = zipfile.ZipFile(io.BytesIO(render_deck_bundle(specs, deck_kind=deck_kind)))
    path = tmp_path / f"{deck_kind}-{len(specs)}.pptx"
    path.write_bytes(bundle.read("deck.pptx"))
    return path, json.loads(bundle.read("manifest.json")), bundle


def assert_closing_last(path: Path, appendix_count: int) -> None:
    """The exported deck: canonical 1-25 unchanged, the appendix, then "Let's talk" - once, last."""
    master = registry.get_master()
    canonical = assembly.canonical_fingerprints(master.pptx_path)
    prints = assembly.canonical_fingerprints(path)
    assert len(prints) == 26 + appendix_count
    assert prints[-1] == canonical[25], "the last slide is the canonical Let’s talk slide, unchanged"
    assert prints.count(canonical[25]) == 1, "Let’s talk appears exactly once"
    assert prints[:25] == canonical[:25], "the other canonical slides keep their order and content"
    assert not set(prints[25:-1]) & set(canonical), "only appendix slides sit between them"
    deck = Presentation(str(path))
    assert any(CLOSING_TEXT in text for text in texts(deck.slides[-1]))
    total = len(deck.slides)
    for position in range(25, total - 1):
        assert f"{position + 1} / {total}" in texts(deck.slides[position]), "appendix pages are numbered by their place"


# --- the closing slide is declared by the master, not by an index --------------------------


def test_the_closing_slide_is_the_manifest_role_and_ends_every_order() -> None:
    master = registry.get_master()
    assert registry.closing_slide(master) == 26
    assert registry.canonical_slide_titles()[registry.closing_slide(master) - 1] == CLOSING_TITLE
    for count in (0, 1, 2, 12):
        order = registry.deck_order(master, count)
        assert len(order) == 26 + count and order[-1] == ("canonical", 26)
        assert [ref for kind, ref in order if kind == "canonical"] == [*range(1, 26), 26]
        assert [ref for kind, ref in order if kind == "appendix"] == list(range(count))
        assert len(set(order)) == len(order), "no slide appears twice"
    assert registry.first_appendix_position(master) == 26


def test_another_closing_slide_or_none_follows_the_manifest(monkeypatch: pytest.MonkeyPatch) -> None:
    master = registry.get_master()
    manifest = registry.load_manifest(master)

    def declaring(*numbers: int):
        changed = copy.deepcopy(manifest)
        for slide in changed["slides"]:
            slide.pop("role", None)
            if slide["number"] in numbers:
                slide["role"] = registry.CLOSING_ROLE
        return lambda _master: changed

    monkeypatch.setattr(registry, "load_manifest", declaring(3))
    assert registry.deck_order(master, 2)[-3:] == [("appendix", 0), ("appendix", 1), ("canonical", 3)]
    monkeypatch.setattr(registry, "load_manifest", declaring())
    assert registry.deck_order(master, 1)[-1] == ("appendix", 0), "without a declared closing slide the appendix ends the deck"
    monkeypatch.setattr(registry, "load_manifest", declaring(25, 26))
    with pytest.raises(registry.MasterDeckError, match="more than one closing slide"):
        registry.deck_order(master, 1)


# --- the exported PPTX ---------------------------------------------------------------------


@pytest.mark.parametrize("appendix_count", [0, 1, 3, None], ids=["no-appendix", "one-slide", "three-slides", "full-appendix"])
def test_exported_v1_pptx_ends_with_lets_talk_for_any_appendix_size(appendix_count: int | None, tmp_path: Path) -> None:
    plan = build_master_plan(approved_row(RICH))
    specs = specs_of(plan)
    appendix_specs = specs[25:-1]
    assert len(appendix_specs) > 3, "the full appendix has several client slides"
    chosen = appendix_specs if appendix_count is None else appendix_specs[:appendix_count]
    deck_specs = [*specs[:25], *chosen, specs[-1]]
    path, manifest, bundle = exported(deck_specs, MASTER_V1, tmp_path)
    assert_closing_last(path, len(chosen))
    total = 26 + len(chosen)
    # PDF and previews are rendered from that PPTX: one page and one preview per slide, same order.
    assert manifest["slideCount"] == total and manifest["appendixSlideCount"] == len(chosen)
    assert manifest["previews"] == [f"preview-{n:03d}.png" for n in range(1, total + 1)]
    assert len(Presentation(io.BytesIO(bundle.read("deck.pptx"))).slides) == total


def test_v1_plan_specs_and_pptx_agree_on_the_order(tmp_path: Path) -> None:
    plan = build_master_plan(approved_row(RICH))
    specs = specs_of(plan)
    assert specs[-1]["title"] == CLOSING_TITLE and specs[-1]["layoutId"] == CANONICAL_LAYOUT
    assert [spec["title"] for spec in specs].count(CLOSING_TITLE) == 1
    appendix_count = plan["appendix"]["slide_count"]
    assert [slide["order"] for slide in plan["slides"][25:-1]] == list(range(26, 26 + appendix_count))
    path, _manifest, _bundle = exported(specs, MASTER_V1, tmp_path)
    assert_closing_last(path, appendix_count)
    # Planning and rendering are deterministic: a regeneration of the same source has the same order.
    again = specs_of(build_master_plan(approved_row(RICH)))
    assert again == specs


def test_exported_v2_pptx_ends_with_lets_talk(tmp_path: Path) -> None:
    snapshot = snapshot_for()
    plan = v2.build_plan(snapshot)
    specs = specs_of(plan)
    assert specs[-1]["title"] == CLOSING_TITLE and [spec["title"] for spec in specs].count(CLOSING_TITLE) == 1
    v2.verify_frozen_job(manifest=v2.generation_manifest(snapshot), snapshot=snapshot, plan_json=plan)
    path, manifest, _bundle = exported(specs, MASTER_V2, tmp_path)
    assert_closing_last(path, plan["appendix"]["slide_count"])
    assert manifest["slideCount"] == len(specs)


def test_the_closing_slide_is_never_duplicated_dropped_or_moved_by_assembly(tmp_path: Path) -> None:
    master = registry.get_master()
    plan = build_master_plan(approved_row(RICH))
    contents = [slide["borekSlide"] for slide in plan["slides"][25:-1]]
    final = tmp_path / "final.pptx"
    assembly.assemble(master, contents, final)
    assert_closing_last(final, len(contents))
    # A deck with the closing slide anywhere but last fails verification.
    legacy = tmp_path / "legacy.pptx"
    assembly.assemble(master, contents, legacy, closing_last=False)
    with pytest.raises(assembly.AssemblyError, match="Canonical slide 26 was changed"):
        assembly.verify_canonical_preserved(master, legacy, expected_total=26 + len(contents))


# --- earlier decks: not reused, but a frozen earlier job still renders what it planned -----


def test_a_v1_deck_generated_in_the_earlier_order_is_generated_again() -> None:
    manifest = generation_manifest(approved_row())
    earlier = {key: value for key, value in manifest.items() if key != "deck_layout"}
    assert manifest["deck_layout"] == registry.DECK_LAYOUT
    assert not same_source(manifest, earlier) and same_source(manifest, generation_manifest(approved_row()))


def test_a_v2_frozen_in_the_earlier_order_is_not_reused_but_its_job_stays_valid(tmp_path: Path) -> None:
    current = snapshot_for()
    earlier = copy.deepcopy(current)
    earlier["master"].pop("deck_layout")
    assert not v2.same_source(v2.generation_manifest(current), v2.generation_manifest(earlier)), "regenerated in the new order"
    # A job frozen before BAP-4 keeps its manifest and plan consistent and renders what it planned.
    manifest, plan = v2.generation_manifest(earlier), v2.build_plan(earlier)
    assert "deck_layout" not in manifest and "deck_layout" not in plan
    v2.verify_frozen_job(manifest=manifest, snapshot=earlier, plan_json=plan)
    specs = specs_of(plan)
    assert specs[25]["title"] == CLOSING_TITLE and plan["appendix"]["first_slide"] == 27
    path, manifest, _bundle = exported(specs, MASTER_V2, tmp_path)
    prints = assembly.canonical_fingerprints(path)
    assert prints[:26] == assembly.canonical_fingerprints(registry.get_master().pptx_path)
    assert manifest["slideCount"] == len(specs)
