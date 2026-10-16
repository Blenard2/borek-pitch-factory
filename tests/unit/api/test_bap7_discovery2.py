"""BAP-7: Discovery 2 - generated after the meeting from the approved Discovery 1, the transcript,
the notes and the confirmed findings; reviewed, edited, approved and versioned; the authoritative
source of Master Presentation V2; and, with Discovery 1 and both Master Presentations, available
to the follow-up email."""

from __future__ import annotations

import copy
import email
import email.policy
import hashlib
import io
import json
from datetime import UTC, datetime, timedelta
from uuid import UUID, uuid4

import httpx
import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app.config import settings
from app.services import post_meeting_discovery as discovery2_service
from app.services.data.memory_store import get_memory_store
from app.services.data.supabase_store import SupabaseDataStore
from services.meeting.discovery2 import CATEGORIES, Discovery2Error, apply_section_edits, validate_discovery2
from tests.unit.api.discovery2_helpers import approve_current_discovery2, generate_discovery2
from tests.unit.api.test_followup_email import confirm_email, export, finalized, generate as generate_email, load as load_email, patch as patch_email
from tests.unit.api.test_master_presentation_v2_api import (
    CLIENT,
    EXCLUDED,
    NOTES,
    OBSERVATION,
    OTHER,
    OWNER,
    client,  # noqa: F401  (pytest fixture)
    confirm,
    download,
    generate_v2,
    get,
    headers,
    ok,
    post,
    prepared,
    review,
    slide_texts,
    stored_versions,
)

ROOT_PATH = "/opportunities/{}/post-meeting-discovery"


def d2(client: TestClient, opportunity_id: str, user: UUID = OWNER) -> dict:
    return ok(get(client, ROOT_PATH.format(opportunity_id), user))


def edit(client: TestClient, opportunity_id: str, current: dict, sections: list[dict], user: UUID = OWNER):
    return client.patch(
        ROOT_PATH.format(opportunity_id),
        headers=headers(user),
        json={"expected_version_id": current["id"], "expected_content_hash": current["content_hash"], "sections": sections},
    )


def approve(client: TestClient, opportunity_id: str, current: dict, user: UUID = OWNER):
    return post(
        client,
        f"{ROOT_PATH.format(opportunity_id)}/approve",
        {"expected_version_id": current["id"], "expected_content_hash": current["content_hash"]},
        user,
    )


def section(document: dict, key: str) -> dict:
    return next(item for item in document["sections"] if item["key"] == key)


def meeting(document: dict, key: str) -> list[tuple[str, str]]:
    return [(item["text"], item["source"]) for item in section(document, key)["items"] if item["origin"] in {"meeting", "owner"}]


def confirmed(client: TestClient, user: UUID = OWNER, *, excluded: dict | None = None, **kwargs) -> dict:
    """Prepared opportunity with confirmed findings (the follow-up excluded) and no Discovery 2 yet."""
    base = prepared(client, user, **kwargs)
    view = ok(get(client, f"/opportunities/{base['opportunity']}/post-meeting-review", user))
    ok(
        post(
            client,
            f"/opportunities/{base['opportunity']}/post-meeting-review/confirm",
            {
                "transcript_id": view["extraction"]["transcript_id"],
                "extraction_generated_at": view["extraction"]["generated_at"],
                "review_fingerprint": view["review_fingerprint"],
                "excluded": {"follow_ups": [EXCLUDED]} if excluded is None else excluded,
            },
            user,
        )
    )
    return base


# ------------------------------------------------------------------ generation


def test_discovery2_combines_discovery1_transcript_notes_and_confirmed_findings_and_leaves_discovery1_unchanged(client: TestClient) -> None:
    base = confirmed(client)
    opportunity_id = base["opportunity"]
    store = get_memory_store()
    discovery1_rows = copy.deepcopy(store.discovery_paper_versions)
    discovery1_paper = copy.deepcopy(store.opportunities[UUID(opportunity_id)]["discovery_paper"])

    before = d2(client, opportunity_id)
    assert (before["state"]["status"], before["current_version"], before["input_blockers"]) == ("not_generated", None, [])
    assert before["state"]["can_generate"] is True and before["document"]["status"] == "not_generated"

    generated = generate_discovery2(client, opportunity_id, headers())
    document = validate_discovery2(generated["document"])
    assert (generated["state"]["status"], generated["current_version"]["version_number"], generated["current_version"]["status"]) == ("draft", 1, "draft")
    assert document["generation"] == {"mode": "fixture", "generator_version": "discovery2:v1"}, "fixture output is never labelled as AI"
    assert [item["key"] for item in document["sections"]] == [
        "context", "requirements", "challenges", "priorities", "opportunities", "discussed_solutions", "decisions", "follow_ups", "open_questions",
    ]

    # Carried forward from the approved Discovery 1, in its own wording.
    approved = ok(get(client, f"/opportunities/{opportunity_id}/discovery-paper/approved"))
    thesis = approved["paper_json"]["analysis"]["research"]["core_thesis"]["text"]
    first = section(document, "context")["items"][0]
    assert (first["text"], first["origin"], first["status"], first["source"]) == (thesis, "discovery1", "carried_forward", "discovery1")
    carried = [item for part in document["sections"] for item in part["items"] if item["origin"] == "discovery1"]
    assert len(carried) >= 4 and all(item["evidence"] is None and item["discovery1_ref"] for item in carried)
    assert section(document, "open_questions")["items"], "open Discovery 1 questions are carried forward"

    # New from the meeting: exactly the confirmed findings, with their verified sources.
    assert meeting(document, "requirements") == [("Quotes must go out within one day.", "transcript")]
    assert meeting(document, "challenges") == [("Pricing data sits in three systems.", "transcript")]
    assert meeting(document, "priorities") == [("Start with the sales team.", "both")]
    assert meeting(document, "opportunities") == [(OBSERVATION, "personal_notes")]
    assert meeting(document, "decisions") == [("Run a pilot with ten quotes.", "transcript")]
    assert meeting(document, "follow_ups") == [], "an excluded finding never reaches Discovery 2"
    assert EXCLUDED not in json.dumps(document) and "trade fair" not in json.dumps(document), "nothing unconfirmed is added"
    assert "5 requirements" not in document["summary"]["meeting"] and "1 requirements" in document["summary"]["meeting"]

    # Provenance: the exact source revisions.
    view = review(client, opportunity_id)
    sources = document["sources"]
    assert sources["discovery1"]["version_id"] == base["approved"]
    assert (sources["transcript"]["transcript_id"], sources["transcript"]["revision"]) == (base["transcript"], view["extraction"]["transcript_revision"])
    assert sources["personal_notes"]["sha256"] == hashlib.sha256(NOTES.encode()).hexdigest()
    assert sources["personal_notes"]["updated_at"] == view["personal_notes"]["updated_at"]
    assert sources["meeting_review"]["confirmed_finding_count"] == 6
    assert sources["basis_hash"] == view["discovery2"]["basis_hash"] == generated["current_version"]["basis_hash"]

    # Discovery 1 is read, never written.
    assert store.discovery_paper_versions == discovery1_rows
    assert store.opportunities[UUID(opportunity_id)]["discovery_paper"] == discovery1_paper
    assert [entry["action"] for entry in store.audit_logs.values()].count("discovery2.generate") == 1


def _without_transcript(client: TestClient) -> str:
    opportunity_id = ok(client.post("/opportunities", headers=headers(), json=CLIENT), 201)["id"]
    ok(post(client, f"/opportunities/{opportunity_id}/discovery-paper/generate"))
    ok(post(client, f"/opportunities/{opportunity_id}/discovery-paper/approve"))
    ok(post(client, f"/opportunities/{opportunity_id}/stage1-outputs/generate"))
    ok(post(client, f"/opportunities/{opportunity_id}/workflow/first-meeting-completed"))
    return opportunity_id


def test_generation_needs_an_approved_discovery1_a_transcript_and_a_current_confirmed_meeting(client: TestClient) -> None:
    def refused(opportunity_id: str, step: str) -> None:
        response = post(client, f"{ROOT_PATH.format(opportunity_id)}/generate")
        assert response.status_code == 409 and "DISCOVERY2_INPUTS_NOT_READY" in response.text and step in response.text, response.text
        assert d2(client, opportunity_id)["versions"] == []

    fresh = ok(client.post("/opportunities", headers=headers(), json=CLIENT), 201)["id"]
    refused(fresh, "approve the Discovery document")
    assert "DISCOVERY_NOT_APPROVED" in d2(client, fresh)["input_blockers"] and d2(client, fresh)["state"]["can_generate"] is False

    refused(_without_transcript(client), "upload the meeting transcript")

    base = prepared(client)
    opportunity_id = base["opportunity"]
    refused(opportunity_id, "confirm the meeting findings")
    confirm(client, opportunity_id, discovery2=False)
    ok(client.put(f"/opportunities/{opportunity_id}/personal-notes", headers=headers(), json={"text": NOTES + "\nDecision: Owner sponsors the pilot."}))
    refused(opportunity_id, "analyse the meeting again")


def test_live_generation_reads_the_meeting_and_keeps_only_grounded_items(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    base = confirmed(client)
    opportunity_id = base["opportunity"]
    confirmed(client, OTHER, body={**CLIENT, "client_name": "Südhafen Logistik"}, meeting="Cara: Requirement: Warehouse scans must post in real time.\n", notes=None, excluded={})
    seen: dict = {}

    def model(system: str, user: str, schema: dict) -> dict:
        seen.update(system=system, user=user, schema=schema)
        return {
            "requirements": [
                {"text": "Quotes need a one-day turnaround", "status": "updated", "evidence": "Quotes must go out within one day.", "discovery1_ref": "no-such-element"},
                {"text": "Integration with SAP", "status": "new", "evidence": "we must integrate with SAP", "discovery1_ref": None},
                {"text": "Warehouse scans in real time", "status": "new", "evidence": "Warehouse scans must post in real time.", "discovery1_ref": None},
            ],
            "challenges": [],
            "priorities": [],
            "opportunities": [],
            "discussed_solutions": [{"text": "A quote drafting assistant", "status": "new", "evidence": "A drafting assistant that prepares the quote from the request.", "discovery1_ref": None}],
            "decisions": [{"text": "Everything is decided", "status": "new", "evidence": "Indeed", "discovery1_ref": None}],
            "follow_ups": [],
            "open_questions": [{"text": "Which three systems hold the pricing data?", "status": "new", "evidence": "Pricing data sits in three systems.", "discovery1_ref": None}],
        }

    monkeypatch.setattr(settings, "AI_EXECUTION_MODE", "live")
    monkeypatch.setattr(discovery2_service, "_live_complete", lambda _opportunity: model)
    document = generate_discovery2(client, opportunity_id, headers())["document"]

    assert document["generation"]["mode"] == "live"
    assert "Do not invent" in seen["system"] and "Discovery 1 is carried forward automatically" in seen["system"]
    for label in ("DISCOVERY 1", "CONFIRMED MEETING FINDINGS", "TRANSCRIPT:", "OWNER PERSONAL NOTES:"):
        assert label in seen["user"]
    assert "Warehouse" not in seen["user"] and EXCLUDED not in seen["user"].split("TRANSCRIPT:")[0], "only this opportunity's sources"
    assert seen["schema"]["properties"]["requirements"]["items"]["required"] == ["text", "status", "evidence", "discovery1_ref"]

    requirements = [item for item in section(document, "requirements")["items"] if item["origin"] == "meeting"]
    assert [(item["text"], item["status"], item["source"], item["discovery1_ref"]) for item in requirements] == [
        ("Quotes need a one-day turnaround", "new", "transcript", None),  # an unknown reference is not kept
    ], "fabricated and other-client items are dropped; the confirmed finding is covered by its evidence"
    assert meeting(document, "discussed_solutions") == [("A quote drafting assistant", "transcript")]
    # The model's unsupported decision is dropped, and the confirmed decision the model left out is added.
    assert meeting(document, "decisions") == [("Run a pilot with ten quotes.", "transcript")]
    assert meeting(document, "opportunities") == [(OBSERVATION, "personal_notes")]
    assert meeting(document, "open_questions") == [("Which three systems hold the pricing data?", "transcript")]
    text = json.dumps(document)
    assert "SAP" not in text and "Warehouse" not in text and "Everything is decided" not in text


def test_a_failed_live_generation_is_visible_and_changes_no_version(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    opportunity_id = confirmed(client)["opportunity"]

    def broken(_opportunity: str):
        def complete(system: str, user: str, schema: dict) -> dict:
            raise Discovery2Error("The AI model is not available.", "DISCOVERY2_GENERATION_FAILED")

        return complete

    monkeypatch.setattr(settings, "AI_EXECUTION_MODE", "live")
    monkeypatch.setattr(discovery2_service, "_live_complete", broken)
    failed = post(client, f"{ROOT_PATH.format(opportunity_id)}/generate")
    assert failed.status_code == 400 and "DISCOVERY2_GENERATION_FAILED" in failed.text
    state = d2(client, opportunity_id)
    assert (state["state"]["status"], state["state"]["error"]["code"], state["versions"]) == ("failed", "DISCOVERY2_GENERATION_FAILED", [])
    assert state["state"]["can_generate"] is True and "DISCOVERY2_MISSING" in review(client, opportunity_id)["readiness"]["blockers"]


def test_an_interrupted_generation_blocks_briefly_and_then_can_be_repeated(client: TestClient) -> None:
    opportunity_id = confirmed(client)["opportunity"]
    row = get_memory_store().opportunities[UUID(opportunity_id)]
    working = copy.deepcopy(d2(client, opportunity_id)["document"])
    working.update({"status": "generating", "generated_at": datetime.now(UTC).isoformat().replace("+00:00", "Z")})
    row["post_meeting_discovery"] = working
    busy = post(client, f"{ROOT_PATH.format(opportunity_id)}/generate")
    assert busy.status_code == 409 and "DISCOVERY2_GENERATION_IN_PROGRESS" in busy.text
    working["generated_at"] = (datetime.now(UTC) - timedelta(hours=1)).isoformat().replace("+00:00", "Z")
    state = d2(client, opportunity_id)["state"]
    assert (state["status"], state["error"]["code"]) == ("failed", "DISCOVERY2_GENERATION_INTERRUPTED")
    assert generate_discovery2(client, opportunity_id, headers())["state"]["status"] == "draft"


# ------------------------------------------------------------------ edit, approve, versions


def test_edit_save_approve_and_an_edit_after_approval_is_a_new_draft(client: TestClient) -> None:
    opportunity_id = confirmed(client)["opportunity"]
    generated = generate_discovery2(client, opportunity_id, headers())
    current, document = generated["current_version"], generated["document"]
    requirement = next(item for item in section(document, "requirements")["items"] if item["origin"] == "meeting")
    decisions = [{"id": item["id"], "text": item["text"]} for item in section(document, "decisions")["items"]]

    saved = ok(
        edit(
            client,
            opportunity_id,
            current,
            [
                {"key": "requirements", "items": [{"id": requirement["id"], "text": "Quotes must go out within one working day."}]},
                {"key": "challenges", "items": []},
                {"key": "decisions", "items": [*decisions, {"text": "The owner sponsors the pilot."}]},
            ],
        )
    )
    # A draft is edited in place.
    assert saved["current_version"]["id"] == current["id"] and saved["current_version"]["content_hash"] != current["content_hash"]
    assert len(saved["versions"]) == 1
    edited = section(saved["document"], "requirements")["items"][0]
    assert (edited["text"], edited["origin"], edited["edited"], edited["evidence"]) == (
        "Quotes must go out within one working day.", "meeting", True, "Quotes must go out within one day.",
    ), "an edited item keeps its origin and evidence and is marked as edited"
    assert section(saved["document"], "challenges")["items"] == []
    added = section(saved["document"], "decisions")["items"][-1]
    assert (added["text"], added["origin"], added["source"], added["status"]) == ("The owner sponsors the pilot.", "owner", "owner", "new")

    # An edit made on an older content is refused and changes nothing.
    stale = edit(client, opportunity_id, current, [{"key": "challenges", "items": [{"text": "Overwrite"}]}])
    assert stale.status_code == 409 and "DISCOVERY2_STALE" in stale.text
    assert d2(client, opportunity_id)["current_version"] == saved["current_version"]

    approved = ok(approve(client, opportunity_id, saved["current_version"]))
    assert (approved["state"]["status"], approved["current_version"]["status"]) == ("approved", "approved")
    assert approved["current_version"]["approved_at"] and approved["current_version"]["approved_by"] == str(OWNER)
    first_version = approved["current_version"]
    frozen = ok(get(client, f"{ROOT_PATH.format(opportunity_id)}/versions/{first_version['id']}"))["document"]

    # The approved version is immutable, in the store as well.
    with pytest.raises(HTTPException) as immutable:
        get_memory_store().update_discovery2_draft(
            version_id=UUID(first_version["id"]), user_id=OWNER, document_json=frozen, content_hash="0" * 64, expected_content_hash=first_version["content_hash"]
        )
    assert immutable.value.status_code == 409

    # Editing after approval creates the next draft; the approved version stays as it was.
    after = ok(edit(client, opportunity_id, first_version, [{"key": "requirements", "items": [{"id": requirement["id"], "text": "Quotes go out the same day."}]}]))
    assert (after["current_version"]["version_number"], after["current_version"]["status"]) == (2, "draft")
    assert [(item["version_number"], item["status"]) for item in after["versions"]] == [(1, "approved"), (2, "draft")]
    assert ok(get(client, f"{ROOT_PATH.format(opportunity_id)}/versions/{first_version['id']}"))["document"] == frozen
    state = review(client, opportunity_id)
    assert state["readiness"]["blockers"] == ["DISCOVERY2_NOT_APPROVED"] and state["discovery2"]["stale_reasons"] == ["NEWER_DRAFT"]
    blocked = generate_v2(client, opportunity_id)
    assert blocked.status_code == 409 and "approve Discovery 2" in blocked.text

    listed = ok(get(client, f"{ROOT_PATH.format(opportunity_id)}/versions"))["versions"]
    assert [(item["version_number"], item["status"], item["generation_mode"]) for item in listed] == [(1, "approved", "fixture"), (2, "draft", "fixture")]
    audit = [entry["action"] for entry in get_memory_store().audit_logs.values()]
    assert audit.count("discovery2.edit") == 2 and audit.count("discovery2.approve") == 1


def test_regeneration_adds_a_version_and_history_is_kept(client: TestClient) -> None:
    opportunity_id = confirmed(client)["opportunity"]
    generate_discovery2(client, opportunity_id, headers())
    approve_current_discovery2(client, opportunity_id, headers())
    again = generate_discovery2(client, opportunity_id, headers())
    assert [(item["version_number"], item["status"]) for item in again["versions"]] == [(1, "approved"), (2, "draft")]
    assert again["state"]["status"] == "draft" and again["state"]["approved_version"]["version_number"] == 1
    assert again["document"]["document_id"] != ok(get(client, f"{ROOT_PATH.format(opportunity_id)}/versions/{again['versions'][0]['id']}"))["document"]["document_id"]


def test_an_approval_needs_meeting_content_and_the_current_meeting_basis(client: TestClient) -> None:
    opportunity_id = confirmed(client)["opportunity"]
    generated = generate_discovery2(client, opportunity_id, headers())
    emptied = ok(edit(client, opportunity_id, generated["current_version"], [{"key": key, "items": [
        {"id": item["id"], "text": item["text"]} for item in section(generated["document"], key)["items"] if item["origin"] == "discovery1"
    ]} for key in CATEGORIES]))
    nothing = approve(client, opportunity_id, emptied["current_version"])
    assert nothing.status_code == 400 and "DISCOVERY2_NO_MEETING_CONTENT" in nothing.text

    regenerated = generate_discovery2(client, opportunity_id, headers())
    # The meeting basis changes (the follow-up is confirmed too): this draft is stale.
    view = review(client, opportunity_id)
    ok(post(client, f"/opportunities/{opportunity_id}/post-meeting-review/confirm", {
        "transcript_id": view["extraction"]["transcript_id"],
        "extraction_generated_at": view["extraction"]["generated_at"],
        "review_fingerprint": view["review_fingerprint"],
        "excluded": {},
    }))
    stale = approve(client, opportunity_id, regenerated["current_version"])
    assert stale.status_code == 409 and "DISCOVERY2_STALE" in stale.text
    assert review(client, opportunity_id)["discovery2"]["status"] == "stale"
    fresh = generate_discovery2(client, opportunity_id, headers())
    assert meeting(fresh["document"], "follow_ups") == [(EXCLUDED, "transcript")]
    ok(approve(client, opportunity_id, fresh["current_version"]))


def test_discovery2_is_private_to_the_owner_and_the_opportunity(client: TestClient) -> None:
    mine = confirmed(client)["opportunity"]
    theirs = confirmed(client, OTHER, body={**CLIENT, "client_name": "Südhafen Logistik"}, meeting="Cara: Requirement: Warehouse scans must post in real time.\n", notes=None, excluded={})["opportunity"]
    current = generate_discovery2(client, mine, headers())["current_version"]
    their_version = generate_discovery2(client, theirs, headers(OTHER))["current_version"]
    for response in (
        get(client, ROOT_PATH.format(mine), OTHER),
        post(client, f"{ROOT_PATH.format(mine)}/generate", user=OTHER),
        edit(client, mine, current, [{"key": "challenges", "items": []}], OTHER),
        approve(client, mine, current, OTHER),
        get(client, f"{ROOT_PATH.format(mine)}/versions/{current['id']}", OTHER),
        get(client, f"{ROOT_PATH.format(mine)}/versions/{current['id']}/download/docx", OTHER),
        get(client, f"{ROOT_PATH.format(theirs)}/versions/{current['id']}", OTHER),  # my version under their opportunity
        get(client, f"{ROOT_PATH.format(mine)}/versions/{their_version['id']}"),  # their version under mine
    ):
        assert response.status_code in (403, 404), response.text
    assert client.get(ROOT_PATH.format(mine)).status_code == 401
    assert "Warehouse" not in json.dumps(d2(client, mine)) and "Quotes must" not in json.dumps(d2(client, theirs, OTHER))


def test_the_discovery_documents_download_as_word_files(client: TestClient) -> None:
    base = confirmed(client)
    opportunity_id = base["opportunity"]
    current = generate_discovery2(client, opportunity_id, headers())["current_version"]
    response = get(client, f"{ROOT_PATH.format(opportunity_id)}/versions/{current['id']}/download/docx")
    assert response.status_code == 200 and response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    word = Document(io.BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert "Discovery 2" in text and "Quotes must go out within one day." in text
    assert "Carried forward from Discovery 1" in text and "New from the meeting" in text
    cells = [cell.text for table in word.tables for row in table.rows for cell in row.cells]
    assert any("deterministic test generator" in cell for cell in cells), "fixture output is labelled as such, never as AI"
    first = get(client, f"/opportunities/{opportunity_id}/discovery-paper/versions/{base['approved']}/download/docx")
    assert first.status_code == 200
    thesis = ok(get(client, f"/opportunities/{opportunity_id}/discovery-paper/approved"))["paper_json"]["analysis"]["research"]["core_thesis"]["text"]
    assert thesis in "\n".join(paragraph.text for paragraph in Document(io.BytesIO(first.content)).paragraphs)


# ------------------------------------------------------------------ Master Presentation V2


def test_v2_is_blocked_without_an_approved_current_discovery2_and_built_from_the_approved_version(client: TestClient) -> None:
    base = confirmed(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    (v1_row,) = stored_versions(presentation_id)
    v1_before = copy.deepcopy(v1_row)

    missing = generate_v2(client, opportunity_id)
    assert missing.status_code == 409 and "generate Discovery 2" in missing.text
    generated = generate_discovery2(client, opportunity_id, headers())
    draft = generate_v2(client, opportunity_id)
    assert draft.status_code == 409 and "approve Discovery 2" in draft.text

    document = generated["document"]
    requirement = next(item for item in section(document, "requirements")["items"] if item["origin"] == "meeting")
    solution = next(item for item in section(document, "discussed_solutions")["items"] if item["origin"] == "meeting")
    saved = ok(edit(client, opportunity_id, generated["current_version"], [
        {"key": "requirements", "items": [{"id": requirement["id"], "text": "Quotes must leave the same working day."}]},
        {"key": "discussed_solutions", "items": []},
    ]))
    approved = ok(approve(client, opportunity_id, saved["current_version"]))["current_version"]

    started = ok(generate_v2(client, opportunity_id))
    version = next(row for row in stored_versions(presentation_id) if str(row["id"]) == started["presentation_version_id"])
    manifest = version["generation_source_manifest"]
    assert manifest["approved_discovery2_version_id"] == approved["id"]
    assert manifest["discovery2_content_hash"] == approved["content_hash"]
    assert manifest["approved_discovery_version_id"] == base["approved"]
    texts = " ".join(text for slide in slide_texts(download(client, presentation_id, started["presentation_version_id"])) for text in slide)
    assert "Quotes must leave the same working day." in texts, "the approved Discovery 2 wording is what V2 shows"
    assert "Quotes must go out within one day." not in texts and solution["text"] not in texts, "removed in Discovery 2, absent from V2"
    assert stored_versions(presentation_id)[0] == v1_before, "Master Presentation V1 is preserved"

    # The frozen snapshot names the Discovery 2 version, and an owner-edited statement is the owner's.
    store = get_memory_store()
    snapshot = next(
        (job.get("result_json") or {}).get("_enqueue", {}).get("master_v2_snapshot")
        for job in store.list_generation_jobs_for_opportunity(UUID(opportunity_id))
        if (job.get("result_json") or {}).get("_enqueue", {}).get("master_v2_snapshot")
    )
    assert snapshot["approved_discovery2"]["version_id"] == approved["id"]
    assert snapshot["findings"]["requirements"] == [{"text": "Quotes must leave the same working day.", "source": "personal_notes"}]

    # A changed meeting basis makes Discovery 2 stale; V2 waits for a new approval.
    view = review(client, opportunity_id)
    ok(post(client, f"/opportunities/{opportunity_id}/post-meeting-review/confirm", {
        "transcript_id": view["extraction"]["transcript_id"],
        "extraction_generated_at": view["extraction"]["generated_at"],
        "review_fingerprint": view["review_fingerprint"],
        "excluded": {},
    }))
    stale = generate_v2(client, opportunity_id)
    assert stale.status_code == 409 and "generate Discovery 2 again" in stale.text
    assert ok(get(client, f"/opportunities/{opportunity_id}/master-presentation/v2"))["blockers"] == ["DISCOVERY2_STALE"]


# ------------------------------------------------------------------ follow-up email


def test_all_four_generated_documents_are_available_to_the_follow_up_email(client: TestClient) -> None:
    base = finalized(client)
    opportunity_id, presentation_id = base["opportunity"], base["presentation"]
    v1_id = str(stored_versions(presentation_id)[0]["id"])
    discovery2_id = d2(client, opportunity_id)["state"]["approved_version"]["version_id"]
    draft = ok(generate_email(client, opportunity_id))["draft"]
    items = {item["key"]: item for item in draft["attachments"]}
    assert list(items) == ["pptx", "pdf", "v1_pptx", "v1_pdf", "discovery1_docx", "discovery2_docx"]
    assert all(item["available"] for item in items.values()), [item for item in items.values() if not item["available"]]
    assert {key: (item["artifact"], item["artifact_version_id"]) for key, item in items.items()} == {
        "pptx": ("master_presentation_v2", base["v2"]),
        "pdf": ("master_presentation_v2", base["v2"]),
        "v1_pptx": ("master_presentation_v1", v1_id),
        "v1_pdf": ("master_presentation_v1", v1_id),
        "discovery1_docx": ("discovery_1", base["approved"]),
        "discovery2_docx": ("discovery_2", discovery2_id),
    }
    assert [item["selected"] for item in items.values()] == [False, True, False, False, False, False]
    for item in items.values():
        downloaded = get(client, item["download_path"])
        assert downloaded.status_code == 200 and len(downloaded.content) == item["size_bytes"], item["key"]
        assert get(client, item["download_path"], OTHER).status_code in (403, 404)

    chosen = ok(patch_email(client, opportunity_id, draft["id"], {"expected_revision": draft["revision"], "attachments": {key: True for key in items}}))["draft"]
    assert all(item["selected"] for item in chosen["attachments"])
    assert load_email(client, opportunity_id)["attachments"] == chosen["attachments"]
    unknown = patch_email(client, opportunity_id, draft["id"], {"expected_revision": chosen["revision"], "attachments": {"zip": True}})
    assert unknown.status_code == 422
    done = ok(confirm_email(client, opportunity_id, chosen))["draft"]
    message = email.message_from_bytes(export(client, opportunity_id, done).content, policy=email.policy.default)
    files = {part.get_filename(): part for part in message.iter_attachments()}
    assert len(files) == 6
    discovery2_file = files[items["discovery2_docx"]["file_name"]]
    assert discovery2_file.get_content_type() == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "Discovery 2" in "\n".join(p.text for p in Document(io.BytesIO(discovery2_file.get_content())).paragraphs)
    assert message["X-Borek-Presentation-Version-Id"] == base["v2"]


def test_a_discovery_document_alone_is_never_announced_as_the_attached_presentation(client: TestClient) -> None:
    opportunity_id = finalized(client)["opportunity"]
    draft = ok(generate_email(client, opportunity_id))["draft"]
    sentence = "Please find attached the presentation"
    assert sentence in draft["lengths"]["medium"]["body"], "the V2 PDF is selected by default"
    only_discovery = ok(patch_email(client, opportunity_id, draft["id"], {"expected_revision": draft["revision"], "attachments": {"pdf": False, "discovery2_docx": True}}))["draft"]
    assert all(sentence not in only_discovery["lengths"][name]["body"] for name in ("short", "medium", "extensive"))
    assert [item["key"] for item in only_discovery["attachments"] if item["selected"]] == ["discovery2_docx"]
    both = ok(patch_email(client, opportunity_id, draft["id"], {"expected_revision": only_discovery["revision"], "attachments": {"pdf": True}}))["draft"]
    assert sentence in both["lengths"]["medium"]["body"]


# ------------------------------------------------------------------ module rules


def test_section_edits_are_validated() -> None:
    document = {
        "schema_version": "1.0", "kind": "post_meeting_discovery", "opportunity_id": str(uuid4()), "document_id": str(uuid4()),
        "status": "ready", "generated_at": "2026-10-16T10:00:00Z", "error": None, "generation": {"mode": "fixture", "generator_version": "x"},
        "title": "Discovery 2", "client_name": "Acme", "sources": None, "summary": {"previous": "Thesis.", "meeting": ""},
        "sections": [{"key": key, "title": key, "items": []} for key in ("context", *CATEGORIES, "open_questions")],
    }
    with pytest.raises(Discovery2Error, match="does not belong"):
        apply_section_edits(document, [{"key": "requirements", "items": [{"id": "elsewhere", "text": "x"}]}])
    with pytest.raises(Discovery2Error, match="Unknown or repeated"):
        apply_section_edits(document, [{"key": "requirements", "items": []}, {"key": "requirements", "items": []}])
    with pytest.raises(Discovery2Error, match="longer than"):
        apply_section_edits(document, [{"key": "requirements", "items": [{"text": "x" * 4001}]}])
    with pytest.raises(Discovery2Error):
        validate_discovery2({**document, "sections": document["sections"][:3]})


# ------------------------------------------------------------------ Supabase store


def test_the_supabase_store_versions_discovery2_and_never_rewrites_an_approved_version(monkeypatch: pytest.MonkeyPatch) -> None:
    tables: dict[str, list[dict]] = {"opportunities": [], "post_meeting_discovery_versions": []}
    opportunity_id = uuid4()
    tables["opportunities"].append({"id": str(opportunity_id), "created_by": str(OWNER), "client_name": "Acme", "opportunity_name": "x", "department": "y"})

    def matches(row: dict, params: dict) -> bool:
        return all(str(row.get(key)) == value.removeprefix("eq.") for key, value in params.items() if isinstance(value, str) and value.startswith("eq."))

    def request(self, method, table, *, json_body=None, params=None, headers=None):
        rows = tables.setdefault(table, [])
        params = params or {}
        if method == "POST":
            if any(row["opportunity_id"] == json_body["opportunity_id"] and row["version_number"] == json_body["version_number"] for row in rows):
                return httpx.Response(409, json={"message": "duplicate key"})
            created = {**copy.deepcopy(json_body), "id": str(uuid4())}
            rows.append(created)
            return httpx.Response(201, json=[copy.deepcopy(created)])
        matched = [row for row in rows if matches(row, params)]
        if method == "GET":
            ordered = sorted(matched, key=lambda row: row.get("version_number", 0), reverse="desc" in params.get("order", ""))
            limit = int(params.get("limit", len(ordered) or 1))
            return httpx.Response(200, json=[copy.deepcopy(row) for row in ordered[:limit]])
        if method == "PATCH":
            for row in matched:
                row.update(copy.deepcopy(json_body))
            return httpx.Response(200, json=[copy.deepcopy(row) for row in matched])
        raise AssertionError(method)

    monkeypatch.setattr(SupabaseDataStore, "_request", request)
    monkeypatch.setattr(SupabaseDataStore, "get_opportunity", lambda self, *, opportunity_id, user_id: {"id": opportunity_id, "created_by": OWNER})
    store = SupabaseDataStore("test-token")
    discovery1 = uuid4()
    first = store.create_discovery2_version(
        opportunity_id=opportunity_id, user_id=OWNER, document_id=uuid4(), document_json={"a": 1}, content_hash="1" * 64, basis_hash="b" * 64, discovery1_version_id=discovery1
    )
    assert (first["version_number"], first["status"], first["discovery1_version_id"]) == (1, "draft", discovery1)
    updated = store.update_discovery2_draft(version_id=first["id"], user_id=OWNER, document_json={"a": 2}, content_hash="2" * 64, expected_content_hash="1" * 64)
    assert updated["document_json"] == {"a": 2}
    with pytest.raises(HTTPException) as stale:
        store.update_discovery2_draft(version_id=first["id"], user_id=OWNER, document_json={"a": 3}, content_hash="3" * 64, expected_content_hash="1" * 64)
    assert stale.value.status_code == 409
    approved = store.approve_discovery2_version(version_id=first["id"], user_id=OWNER, expected_content_hash="2" * 64)
    assert approved["status"] == "approved" and approved["approved_by"] == OWNER and approved["approved_at"] is not None
    with pytest.raises(HTTPException) as immutable:
        store.update_discovery2_draft(version_id=first["id"], user_id=OWNER, document_json={"a": 4}, content_hash="4" * 64, expected_content_hash="2" * 64)
    assert immutable.value.status_code == 409 and tables["post_meeting_discovery_versions"][0]["document_json"] == {"a": 2}
    second = store.create_discovery2_version(
        opportunity_id=opportunity_id, user_id=OWNER, document_id=uuid4(), document_json={"b": 1}, content_hash="5" * 64, basis_hash="b" * 64, discovery1_version_id=discovery1
    )
    assert second["version_number"] == 2
    assert [row["version_number"] for row in store.list_discovery2_versions(opportunity_id=opportunity_id, user_id=OWNER)] == [1, 2]
    # Two requests that computed the same next number: the unique index lets one through.
    monkeypatch.setattr(store, "_request", lambda method, table, **kwargs: request(store, method, table, **kwargs) if method != "GET" else httpx.Response(200, json=[{"version_number": 1}]))
    with pytest.raises(HTTPException) as raced:
        store.create_discovery2_version(
            opportunity_id=opportunity_id, user_id=OWNER, document_id=uuid4(), document_json={}, content_hash="6" * 64, basis_hash="b" * 64, discovery1_version_id=discovery1
        )
    assert raced.value.status_code == 409 and "DISCOVERY2_CONCURRENT_UPDATE" in str(raced.value.detail)


def test_the_migration_keeps_approved_versions_immutable_and_rows_owned() -> None:
    from pathlib import Path

    sql = (Path(__file__).resolve().parents[3] / "apps/services/api/supabase/migrations/044_bap7_post_meeting_discovery.sql").read_text(encoding="utf-8")
    assert "ADD COLUMN IF NOT EXISTS post_meeting_discovery JSONB" in sql
    assert "CREATE TABLE IF NOT EXISTS public.post_meeting_discovery_versions" in sql
    assert "ENABLE ROW LEVEL SECURITY" in sql and "WHERE created_by = auth.uid()" in sql
    assert "Approved Discovery 2 versions are immutable" in sql and "BEFORE UPDATE OR DELETE" in sql
    assert "post_meeting_discovery_versions_opportunity_version_key" in sql and "REFERENCES public.discovery_paper_versions(id)" in sql
